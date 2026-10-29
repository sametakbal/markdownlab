"""Smoke-test the packaged executable: convert a DOCX with Turkish text.

    python scripts/smoke_test.py dist/MarkdownLab/MarkdownLab.exe

The exe is a windowed (GUI subsystem) app, so it is run via subprocess, which
waits for it to exit (PowerShell would not).
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import docx

HEADING = "Başlık Şİğ ıüöç"


def main(exe: str) -> int:
    # CreateProcess does not resolve relative paths with forward slashes
    exe = str(Path(exe).resolve())
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "örnek.docx"
        out = Path(tmp) / "out"
        d = docx.Document()
        d.add_heading(HEADING, 1)
        d.add_paragraph().add_run("kalın").bold = True
        d.save(src)

        proc = subprocess.run([exe, "--lang", "en", "-o", str(out), str(src)], timeout=300)
        md_path = out / "örnek.md"
        if proc.returncode != 0 or not md_path.is_file():
            print(f"exe failed (exit {proc.returncode}), no output at {md_path}", file=sys.stderr)
            return 1
        md = md_path.read_text(encoding="utf-8")
        if f"# {HEADING}" not in md or "**kalın**" not in md:
            print(f"unexpected output:\n{md}", file=sys.stderr)
            return 1
    print("smoke test ok")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
