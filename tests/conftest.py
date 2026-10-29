from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytest

TR_TEXT = "Iğdır'da şişeler, İSTANBUL ığ — Çağrı Öğütücü ü"
TR_UPPER = "ÇĞİÖŞÜ çğıöşü âîû"

FONT_DIR = Path("C:/Windows/Fonts")


@pytest.fixture(autouse=True)
def _turkish_messages():
    from markdownlab import i18n

    i18n.set_language("tr")
    yield
    i18n.set_language(i18n.DEFAULT_LANGUAGE)


def _png_bytes(color=(200, 30, 30)) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), False)
    pix.set_rect(pix.irect, color)
    return pix.tobytes("png")


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory) -> Path:
    docx = pytest.importorskip("docx")
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Ana Başlık Şeması", level=1)
    p = d.add_paragraph("Normal metin: ")
    p.add_run(TR_TEXT)
    p = d.add_paragraph()
    p.add_run("kalın ğüş").bold = True
    p.add_run(" ve ")
    p.add_run("italik İıÖ").italic = True
    p.add_run(" ve ")
    p.add_run("altı çizili").underline = True
    p.add_run(" ve ")
    p.add_run("üstü çizili").font.strike = True
    d.add_heading("İkinci Düzey Başlık", level=2)
    d.add_paragraph("Birinci madde çay", style="List Bullet")
    d.add_paragraph("İkinci madde şeker", style="List Bullet")
    d.add_paragraph("Numaralı öğe", style="List Number")
    d.add_paragraph("Özel karakterler: 5 * 3 = 15, a_b_c, [köşeli], <etiket>, x|y")
    d.add_paragraph(TR_UPPER)

    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    data = [["Ad", "Şehir", "Not"], ["Ayşe", "Iğdır", "A|B"], ["Çağrı", "İzmir", "ğ"]]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            t.cell(r, c).text = val

    t2 = d.add_table(rows=2, cols=2)
    t2.style = "Table Grid"
    merged = t2.cell(0, 0).merge(t2.cell(0, 1))
    merged.text = "Birleşik başlık"
    t2.cell(1, 0).text = "sol"
    t2.cell(1, 1).text = "sağ"

    d.add_picture(io.BytesIO(_png_bytes()), width=Inches(1))
    path = tmp_path_factory.mktemp("docx") / "örnek belge.docx"
    d.save(path)
    return path


PDF_HTML = f"""
<h1>PDF Başlığı Ğüşİ</h1>
<p>{TR_TEXT}</p>
<p>Bu cümlede <b>kalın şekil</b> ve <i>eğik yazı</i> var.</p>
<h2>Alt Bölüm Çalışması</h2>
<p>{TR_UPPER}</p>
<ul><li>birinci öğe</li><li>ikinci öğe</li></ul>
<table border="1" style="border-collapse: collapse">
<tr><th>Ad</th><th>Şehir</th></tr>
<tr><td>Ayşe</td><td>Iğdır</td></tr>
<tr><td>Çağrı</td><td>İzmir</td></tr>
</table>
"""


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory) -> Path:
    if not (FONT_DIR / "arial.ttf").exists():
        pytest.skip("Türkçe glifli test fontu (Arial) yok")
    css = """
    @font-face {font-family: T; src: url(arial.ttf);}
    @font-face {font-family: T; src: url(arialbd.ttf); font-weight: bold;}
    @font-face {font-family: T; src: url(ariali.ttf); font-style: italic;}
    * {font-family: T;}
    td, th {padding: 4px; border: 1px solid black;}
    """
    story = pymupdf.Story(PDF_HTML, user_css=css, archive=pymupdf.Archive(str(FONT_DIR)))
    path = tmp_path_factory.mktemp("pdf") / "örnek.pdf"
    writer = pymupdf.DocumentWriter(str(path))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()

    doc = pymupdf.open(path)
    doc[0].insert_image(pymupdf.Rect(400, 700, 500, 775), stream=_png_bytes((20, 90, 200)))
    doc.saveIncr()
    doc.close()
    return path


@pytest.fixture(scope="session")
def scanned_pdf(sample_pdf, tmp_path_factory) -> Path:
    """Image-only PDF without a text layer (simulates a scanned document)."""
    src = pymupdf.open(sample_pdf)
    out = pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=200)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, pixmap=pix)
    path = tmp_path_factory.mktemp("scan") / "taranmış.pdf"
    out.save(path)
    return path
